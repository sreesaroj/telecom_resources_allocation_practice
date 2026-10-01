"""Idempotently import the two client master workbooks into an initialized GSUS DB.

This runs inside ``odoo-bin shell`` and expects ``env``. Placeholder people are
converted to gaps, never employees. Source anomalies remain visible in remarks.
"""
import re
import hashlib
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
import openpyxl

VIL_PATH = Path(r"C:\Users\HP\Downloads\VIL Manpower Details.xlsx")
TEAM_PATH = Path(r"C:\Users\HP\Downloads\FRT team List.xlsx")

Partner = env["res.partner"]
Project = env["project.project"]
Zone = env["gsus.zone"]
District = env["gsus.district"]
Location = env["gsus.location"]
Resource = env["gsus.resource"]
ResourceType = env["gsus.resource.type"]
Requirement = env["gsus.requirement"]
Gap = env["gsus.gap"]
Employee = env["hr.employee"]
Vehicle = env["fleet.vehicle"]

types = {record.code: record for record in ResourceType.search([])}
shifts = {record.code: record for record in env["gsus.shift"].search([])}
stats = Counter()


def clean(value):
    return re.sub(r"\s+", " ", str(value or "").replace("�", " ")).strip()


def slug(value):
    return re.sub(r"[^A-Z0-9]+", "-", clean(value).upper()).strip("-")[:35]


LOCATION_NORMALIZATION = {
    "purasaivakkam": "Purasaiwakkam", "tiruvannamalai": "Thiruvannamalai",
    "dindivanam": "Tindivanam", "adayar 12 hrs": "Adayar",
    "salem 12 hrs": "Salem", "salem 24 hrs": "Salem",
    "chennai ekkaduthangal night": "Ekkaduthangal", "perungudi night": "Perungudi",
}


def location_name(value):
    value = clean(value)
    normalized_key = re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()
    return LOCATION_NORMALIZATION.get(normalized_key, value)


def get_customer(name):
    record = Partner.search([("name", "=", name), ("is_telecom_customer", "=", True)], limit=1)
    return record or Partner.create({"name": name, "is_company": True, "is_telecom_customer": True})


def get_project(name, customer, project_type="frt_maintenance"):
    record = Project.search([("name", "=", name), ("customer_id", "=", customer.id)], limit=1)
    return record or Project.create({"name": name, "customer_id": customer.id, "telecom_project_type": project_type, "telecom_active": True})


def get_zone(name, project, prefix, reuse=False):
    record = Zone.search([("name", "=", name), ("project_id", "=", project.id)], limit=1)
    if not record and reuse:
        record = Zone.search([("name", "=", name), ("project_id", "=", False)], limit=1)
        if record: record.project_id = project
    if not record:
        code = f"{prefix}-{slug(name)}"[:40]
        record = Zone.search([("code", "=", code)], limit=1) or Zone.create({"name": name, "code": code, "project_id": project.id})
    return record


def get_district(name, zone, prefix):
    name = clean(name)
    if not name: return District
    record = District.search([("name", "=", name), ("zone_id", "=", zone.id)], limit=1)
    code = f"{prefix}-{slug(name)}"[:40]
    return record or District.search([("code", "=", code)], limit=1) or District.create({"name": name, "code": code, "zone_id": zone.id, "circle": "TN"})


def get_location(name, zone, prefix, hours=None, district=None, reuse=False):
    name = location_name(name)
    record = Location.search([("name", "=", name), ("zone_id", "=", zone.id)], limit=1)
    if not record and reuse:
        record = Location.search([("name", "=", name), ("zone_id", "=", zone.id)], limit=1)
    if not record:
        code = f"{prefix}-{slug(name)}"[:40]
        record = Location.search([("code", "=", code)], limit=1) or Location.create({"name": name, "code": code, "zone_id": zone.id})
    values = {"circle": "TN"}
    if hours: values["frt_hours"] = hours
    if district: values["district_id"] = district.id
    record.write(values)
    return record


def set_requirement(location, shift, resource_type, qty, note):
    record = Requirement.search([("location_id", "=", location.id), ("shift_id", "=", shift.id), ("resource_type_id", "=", resource_type.id)], limit=1)
    values = {"location_id": location.id, "shift_id": shift.id, "resource_type_id": resource_type.id, "required_count": qty, "notes": note}
    if record: record.write(values)
    else: Requirement.create(values); stats["requirements"] += 1


def get_employee(name, code, phone=None):
    raw_barcode = re.sub(r"[^A-Z0-9]", "", clean(code).upper())
    barcode = raw_barcode if raw_barcode and len(raw_barcode) <= 18 else "GS" + hashlib.sha1(clean(code).encode("utf-8")).hexdigest()[:16].upper()
    employee = Employee.search([("barcode", "=", barcode)], limit=1)
    if not employee:
        employee = Employee.create({"name": clean(name), "barcode": barcode, "work_phone": clean(phone) or False})
        stats["employees"] += 1
    return employee


def add_person(name, code, resource_type, customer, project, zone, location, shift, ownership="not_specified", vendor=None, old_id=None, new_id=None, phone=None, note=None):
    if clean(name).lower() in ("vacant", "not available", "-", "n/a", ""): return False
    employee = get_employee(name, code, phone)
    record = Resource.search([("code", "=", code)], limit=1)
    values = {"name": clean(name), "code": code, "resource_type_id": resource_type.id, "employee_id": employee.id,
        "old_employee_id": old_id or False, "new_employee_id": new_id or False, "ownership": ownership,
        "vendor_id": vendor.id if vendor else False, "customer_id": customer.id, "project_id": project.id,
        "zone_id": zone.id, "location_id": location.id, "shift_id": shift.id, "status": "available", "remarks": note or False}
    if record: record.write(values)
    else: Resource.create(values); stats["people_resources"] += 1
    return True


def add_vehicle(plate, code, customer, project, zone, location, shift, used_plates):
    plate = clean(plate).upper()
    brand = env["fleet.vehicle.model.brand"].search([("name", "=", "GSUS Imported")], limit=1) or env["fleet.vehicle.model.brand"].create({"name": "GSUS Imported"})
    model = env["fleet.vehicle.model"].search([("name", "=", "FRT Vehicle"), ("brand_id", "=", brand.id)], limit=1) or env["fleet.vehicle.model"].create({"name": "FRT Vehicle", "brand_id": brand.id})
    vehicle = Vehicle.search([("license_plate", "=", plate)], limit=1) or Vehicle.create({"model_id": model.id, "license_plate": plate})
    conflict = plate in used_plates and used_plates[plate] != location.display_name
    record = Resource.search([("code", "=", code)], limit=1)
    note = f"SOURCE CONFLICT: vehicle also listed at {used_plates[plate]}. Verify assignment dates." if conflict else "Imported from VIL manpower roster."
    values = {"name": f"{location.name} FRT Vehicle", "code": code, "resource_type_id": types["VEHICLE"].id, "vehicle_id": vehicle.id,
        "customer_id": customer.id, "project_id": project.id, "zone_id": zone.id, "location_id": location.id,
        "shift_id": shift.id, "status": "unavailable" if conflict else "available", "ownership": "not_specified", "remarks": note}
    if record: record.write(values)
    else: Resource.create(values); stats["vehicle_resources"] += 1
    if conflict: stats["vehicle_conflicts"] += 1
    else: used_plates[plate] = location.display_name


def add_gap(location, shift, resource_type, qty, reason):
    if qty <= 0: return
    record = Gap.search([("location_id", "=", location.id), ("shift_id", "=", shift.id), ("resource_type_id", "=", resource_type.id), ("reason", "=", reason), ("state", "!=", "resolved")], limit=1)
    if record: record.gap_count = qty
    else: Gap.create({"location_id": location.id, "shift_id": shift.id, "resource_type_id": resource_type.id, "gap_count": qty, "gap_start_date": date(2026, 8, 24), "reason": reason}); stats["gaps"] += 1


# Workbook 1: VIL manpower, vehicle and shift rosters.
vi_customer = get_customer("Vodafone Idea (VI)")
oam = get_project("VI O&M", vi_customer, "frt_maintenance")
mobility = get_project("VI Mobility", vi_customer, "mobility")
ws = openpyxl.load_workbook(VIL_PATH, data_only=True)["FRT  Report"]
used_plates = {}

for start, end, project, prefix, mode in [(6, 13, oam, "VIOAM", "24"), (18, 27, oam, "VIOAM", "12"), (32, 44, mobility, "VIMOB", "mobility")]:
    for row_no in range(start, end + 1):
        row = [ws.cell(row_no, col).value for col in range(2, 16)]
        hours = "24" if "24" in clean(row[2]) else "12"
        zone = get_zone(clean(row[4]), project, prefix, reuse=(project == oam))
        location = get_location(row[5], zone, prefix, hours=hours, reuse=(project == oam))
        vehicle_shift = shifts["FULL24"] if hours == "24" else shifts["DAY12"]
        add_vehicle(row[3], f"{prefix}-R{row_no}-VEH-{slug(location.name)}-{vehicle_shift.code}", vi_customer, project, zone, location, vehicle_shift, used_plates)
        set_requirement(location, vehicle_shift, types["VEHICLE"], 1, "Vehicle scope from VIL manpower roster")
        if mode == "24":
            slots = [(row[6], "DRIVER", "DAY12", 1), (row[7], "HELPER", "DAY12", 1), (row[8], "HELPER", "DAY12", 2),
                     (row[9], "DRIVER", "NIGHT12", 1), (row[10], "SPLICER", "NIGHT12", 1), (row[11], "HELPER", "NIGHT12", 1), (row[12], "HELPER", "NIGHT12", 2)]
            for code, shift_code, qty in [("DRIVER", "DAY12", 1), ("HELPER", "DAY12", 2), ("DRIVER", "NIGHT12", 1), ("SPLICER", "NIGHT12", 1), ("HELPER", "NIGHT12", 2)]:
                set_requirement(location, shifts[shift_code], types[code], qty, "Shift scope from VIL manpower roster")
        elif mode == "12":
            slots = [(row[6], "DRIVER", "DAY12", 1), (row[7], "HELPER", "DAY12", 1), (row[8], "HELPER", "DAY12", 2)]
            set_requirement(location, shifts["DAY12"], types["DRIVER"], 1, "Scope from VIL ROTN roster")
            set_requirement(location, shifts["DAY12"], types["HELPER"], 2, "Scope from VIL ROTN roster")
        else:
            names = [x.strip() for x in clean(row[6]).split("/") if x.strip()]
            slots = [(name, "DRIVER", "DAY12" if i == 0 else "NIGHT12", i + 1) for i, name in enumerate(names)]
            for _, code, shift_code, _ in slots: set_requirement(location, shifts[shift_code], types[code], 1, "Scope from VI Mobility roster")
        missing = defaultdict(int)
        for person_name, type_code, shift_code, slot_no in slots:
            code = f"{prefix}-R{row_no}-{slug(location.name)}-{type_code[:3]}-{shift_code}-{slot_no}"
            if not add_person(person_name, code, types[type_code], vi_customer, project, zone, location, shifts[shift_code], note="Imported from VIL shift roster."):
                missing[(type_code, shift_code)] += 1
        for (type_code, shift_code), qty in missing.items(): add_gap(location, shifts[shift_code], types[type_code], qty, "Vacant position imported from VIL manpower roster")


# Workbook 2: detailed employee/FRT ownership master. Customer is explicitly
# marked unconfirmed because the workbook does not identify one.
team_customer = get_customer("Customer To Be Confirmed – FRT Team List")
team_project = get_project("Southern TN FRT Team Operations", team_customer, "frt_maintenance")
vendor = Partner.search([("name", "=", "External Manpower Vendor")], limit=1) or Partner.create({"name": "External Manpower Vendor", "is_company": True})
ws = openpyxl.load_workbook(TEAM_PATH, data_only=True)["Data"]
team_rows = list(ws.iter_rows(min_row=2, values_only=True))
by_frt = defaultdict(list)
for row in team_rows: by_frt[clean(row[11])].append(row)
for frt_code, rows in by_frt.items():
    first = rows[0]; zone = get_zone(clean(first[1]), team_project, "TNFRT")
    district = get_district(first[6], zone, "TNFRT")
    location = get_location(frt_code, zone, "TNFRT", hours="12", district=district)
    for type_code, qty in [("SPLICER", 1), ("DRIVER", 1), ("HELPER", 2)]: set_requirement(location, shifts["DAY12"], types[type_code], qty, "Standard four-person FRT scope inferred from complete teams in FRT Team List")
    actual = Counter()
    for row in rows:
        designation = clean(row[9]); type_code = "HELPER" if designation.lower().startswith("helper") else designation.upper()
        person_name = clean(row[8]); old_id = clean(row[3]); new_id = clean(row[4]); phone = clean(row[12]); category = clean(row[10]).lower()
        ownership = "vendor" if category.startswith("vendor") else "own_bike" if "bike" in category else "gsus_own"
        identity = new_id if new_id not in ("", "-") else old_id if old_id not in ("", "-") else f"TNFRT-{slug(frt_code)}-{type_code[:3]}-{len(actual)+1}"
        created = add_person(person_name, identity, types[type_code], team_customer, team_project, zone, location, shifts["DAY12"], ownership=ownership, vendor=vendor if ownership == "vendor" else None, old_id=old_id if old_id != "-" else None, new_id=new_id if new_id != "-" else None, phone=phone if phone.lower() != "n/a" else None, note=f"Source category: {clean(row[10])}")
        if created: actual[type_code] += 1
    for type_code, required in [("SPLICER", 1), ("DRIVER", 1), ("HELPER", 2)]:
        add_gap(location, shifts["DAY12"], types[type_code], max(required - actual[type_code], 0), "Missing position inferred from FRT Team List")

env.cr.commit()
print("IMPORT_STATS", dict(stats))
print("MASTER_COUNTS", {model: env[model].search_count([]) for model in ["res.partner", "project.project", "gsus.zone", "gsus.district", "gsus.location", "hr.employee", "fleet.vehicle", "gsus.resource", "gsus.requirement", "gsus.gap"]})
