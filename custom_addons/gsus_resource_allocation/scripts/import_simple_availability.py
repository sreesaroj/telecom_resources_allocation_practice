"""Import Sheet1 into the simplified current FRT availability screen.

Run through Odoo shell with a global ``env`` and pass the workbook path as
``GSUS_SIMPLE_XLSX`` in the process environment.
"""
import os
import openpyxl


path = os.environ["GSUS_SIMPLE_XLSX"]
book = openpyxl.load_workbook(path, data_only=True)
sheet = book["Sheet1"]

type_map = {
    "Vehicle": (4, 5, 7, None),
    "Helper": (8, 9, 12, 11),
    "Patroller": (13, 14, 16, None),
    "Driver": (17, 18, 21, 20),
}
ResourceType = env["gsus.resource.type"]
Availability = env["gsus.frt.availability"]
Line = env["gsus.frt.availability.line"]
Resource = env["gsus.resource"]

created_headers = created_lines = 0
for row in range(4, sheet.max_row + 1):
    zone_name = sheet.cell(row, 2).value
    frt_name = sheet.cell(row, 3).value
    if not zone_name or not frt_name:
        continue
    location = env["gsus.location"].search([
        ("name", "=ilike", str(frt_name).strip()),
        ("zone_id.name", "=ilike", str(zone_name).strip()),
    ], limit=1)
    if not location:
        continue
    header = Availability.search([("location_id", "=", location.id)], limit=1)
    values = {
        "customer_id": location.customer_id.id,
        "project_id": location.project_id.id,
        "zone_id": location.zone_id.id,
        "location_id": location.id,
        "notes": "Imported from GSUS Resources availability Sheet1",
    }
    if not values["customer_id"] or not values["project_id"]:
        continue
    if header:
        header.write(values)
    else:
        header = Availability.create(values)
        created_headers += 1
    for type_name, (scope_col, available_col, remarks_col, alternate_col) in type_map.items():
        resource_type = ResourceType.search([("name", "=ilike", type_name)], limit=1)
        if not resource_type:
            continue
        scope = int(sheet.cell(row, scope_col).value or 0)
        available = int(sheet.cell(row, available_col).value or 0)
        alternate = int(sheet.cell(row, alternate_col).value or 0) if alternate_col else 0
        remarks = sheet.cell(row, remarks_col).value
        resources = Resource.search([
            ("location_id", "=", location.id),
            ("resource_type_id", "=", resource_type.id),
            ("status", "=", "available"),
        ], limit=available)
        vals = {
            "availability_id": header.id,
            "resource_type_id": resource_type.id,
            "scope_count": scope,
            "available_count": available,
            "alternate_count": alternate,
            "available_resource_ids": [(6, 0, resources.ids)],
            "remarks": str(remarks) if remarks else False,
        }
        line = Line.search([
            ("availability_id", "=", header.id),
            ("resource_type_id", "=", resource_type.id),
        ], limit=1)
        if line:
            line.write(vals)
        else:
            Line.create(vals)
            created_lines += 1

env.cr.commit()
print("Imported simplified availability:", created_headers, "FRTs and", created_lines, "scope lines")
