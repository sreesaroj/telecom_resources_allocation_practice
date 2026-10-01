"""Reset GSUS business data and build the 8077 Vodafone client demo.

This intentionally preserves Odoo users/configuration and GSUS resource-type
masters. Run only against the dedicated demonstration database.
"""
import os
import re
import openpyxl


if env.cr.dbname != "gsus_resource_testing":
    raise RuntimeError("Safety stop: this reset script may run only on gsus_resource_testing")

path = os.environ["GSUS_DEMO_XLSX"]
book = openpyxl.load_workbook(path, data_only=True)
sheet1 = book["Sheet1"]
sheet2 = book["Sheet2"]


def slug(value):
    return re.sub(r"[^A-Z0-9]+", "-", str(value).upper()).strip("-")


# Remove only GSUS/demo business data. Master resource types and Odoo users stay.
env["res.partner"].search([]).write({"allowed_project_ids": [(5, 0, 0)]})
old_resources = env["gsus.resource"].search([])
old_employees = old_resources.mapped("employee_id")
env["hr.attendance"].search([("employee_id", "in", old_employees.ids)]).unlink()
env["gsus.resource.request"].search([]).unlink()
env["gsus.frt.availability"].search([]).unlink()
env["gsus.gap"].search([]).unlink()
env["gsus.availability"].search([]).unlink()
env["gsus.assignment"].search([]).unlink()
env["gsus.requirement"].search([]).unlink()
env["gsus.resource"].search([]).unlink()
old_employees.exists().unlink()
env["gsus.worker"].search([]).unlink()
env["fleet.vehicle"].search([("gsus_customer_id", "!=", False)]).unlink()
env["gsus.location"].search([]).unlink()
env["gsus.district"].search([]).unlink()
env["gsus.zone"].search([]).unlink()
env["gsus.resource.type"].search([]).write({"customer_id": False, "project_id": False})
projects = env["project.project"].search([("telecom_active", "=", True)])
projects.unlink()
customers = env["res.partner"].search([("is_telecom_customer", "=", True)])
customers.unlink()

customer = env["res.partner"].create({
    "name": "Vodafone Idea Limited",
    "is_company": True,
    "is_telecom_customer": True,
    "phone": "+91 44 4000 2026",
    "email": "operations.demo@vodafoneidea.com",
})
project = env["project.project"].create({
    "name": "VI Tamil Nadu FRT Operations 2026",
    "partner_id": customer.id,
    "customer_id": customer.id,
    "telecom_project_type": "frt_maintenance",
    "telecom_active": True,
})

resource_types = {}
for name, code, category in [
    ("Vehicle", "VEHICLE", "vehicle"), ("Helper", "HELPER", "person"),
    ("Patroller", "PATROLLER", "person"), ("Driver", "DRIVER", "person"),
    ("Splicer", "SPLICER", "person"), ("Engineer (OME)", "ENGINEER", "person"),
]:
    record = env["gsus.resource.type"].search([("code", "=", code)], limit=1)
    if not record:
        record = env["gsus.resource.type"].create({"name": name, "code": code, "category": category})
    record.write({"customer_id": customer.id, "project_id": project.id})
    resource_types[name] = record

zone_names = []
for row in range(4, sheet1.max_row + 1):
    value = sheet1.cell(row, 2).value
    if value and value not in zone_names:
        zone_names.append(value)
zones = {
    name: env["gsus.zone"].create({
        "name": name, "code": slug(name), "project_id": project.id,
    }) for name in zone_names
}

locations = {}
availability_by_frt = {}
columns = {
    "Vehicle": (4, 5, 7, None),
    "Helper": (8, 9, 12, 11),
    "Patroller": (13, 14, 16, None),
    "Driver": (17, 18, 21, 20),
}
person_sequence = 1
vehicle_sequence = 1
day_shift = env.ref("gsus_resource_allocation.shift_day_12")
night_shift = env.ref("gsus_resource_allocation.shift_night_12")

for row in range(4, sheet1.max_row + 1):
    zone_name = sheet1.cell(row, 2).value
    frt_name = sheet1.cell(row, 3).value
    if not zone_name or not frt_name:
        continue
    location = env["gsus.location"].create({
        "name": frt_name,
        "code": f"FRT-{row - 3:02d}-{slug(frt_name)[:12]}",
        "zone_id": zones[zone_name].id,
        "responsible_id": env.user.id,
    })
    locations[frt_name] = location
    availability = env["gsus.frt.availability"].create({
        "customer_id": customer.id,
        "project_id": project.id,
        "zone_id": zones[zone_name].id,
        "location_id": location.id,
        "responsible_id": env.user.id,
        "notes": "Client demonstration data based on Sheet1 and Sheet2.",
    })
    availability_by_frt[frt_name] = availability

    for type_name, (scope_col, available_col, remarks_col, alternate_col) in columns.items():
        scope = int(sheet1.cell(row, scope_col).value or 0)
        available = int(sheet1.cell(row, available_col).value or 0)
        alternate = int(sheet1.cell(row, alternate_col).value or 0) if alternate_col else 0
        remarks = sheet1.cell(row, remarks_col).value
        rtype = resource_types[type_name]
        requirement = env["gsus.requirement"].create({
            "location_id": location.id,
            "resource_type_id": rtype.id,
            "required_count": scope,
            "notes": "Scope imported from Sheet1",
        })

        persons = env["gsus.resource"]
        vehicles = env["fleet.vehicle"]
        if type_name == "Vehicle":
            for number in range(available):
                vehicle = env["fleet.vehicle"].create({
                    "license_plate": f"TN-DEMO-{vehicle_sequence:03d}",
                    "gsus_make_model": "FRT Utility Vehicle",
                    "gsus_customer_id": customer.id,
                    "gsus_project_id": project.id,
                    "gsus_zone_id": zones[zone_name].id,
                    "gsus_location_id": location.id,
                    "gsus_status": "available",
                    "insurance_policy_no": f"VI-DEMO-POL-{vehicle_sequence:03d}",
                    "insurance_company": "Demo General Insurance",
                    "insurance_expiry_date": "2027-08-24",
                    "permit_expiry_date": "2027-12-31",
                    "pollution_expiry_date": "2027-02-28",
                })
                vehicles |= vehicle
                vehicle_sequence += 1
        else:
            for number in range(available):
                person = env["gsus.resource"].create({
                    "name": f"{frt_name} {type_name} {number + 1:02d}",
                    "code": f"VI-DEMO-{person_sequence:04d}",
                    "resource_type_id": rtype.id,
                    "ownership": "gsus_own",
                    "customer_id": customer.id,
                    "project_id": project.id,
                    "zone_id": zones[zone_name].id,
                    "location_id": location.id,
                    "shift_id": day_shift.id if person_sequence % 2 else night_shift.id,
                    "status": "available",
                    "remarks": "Demo person generated from Sheet1 availability count",
                })
                persons |= person
                person_sequence += 1

        requirement.write({
            "assigned_resource_ids": [(6, 0, persons.ids)],
            "assigned_vehicle_ids": [(6, 0, vehicles.ids)],
        })

        env["gsus.frt.availability.line"].create({
            "availability_id": availability.id,
            "resource_type_id": rtype.id,
            "scope_count": scope,
            "available_count": available,
            "alternate_count": alternate,
            "available_resource_ids": [(6, 0, persons.ids)],
            "available_vehicle_ids": [(6, 0, vehicles.ids)],
            "remarks": str(remarks) if remarks else False,
        })

# Add the Engineer/Splicer gaps present only in Sheet2 without overwriting
# Sheet1's four primary resource columns.
extra_gaps = {}
for row in range(12, sheet2.max_row + 1):
    description = str(sheet2.cell(row, 3).value or "").strip()
    frt_name = str(sheet2.cell(row, 4).value or "").strip()
    gap = int(sheet2.cell(row, 6).value or 0)
    if description not in ("Engineer", "Splicer") or frt_name not in locations or not gap:
        continue
    type_name = "Engineer (OME)" if description == "Engineer" else "Splicer"
    extra_gaps[(frt_name, type_name)] = extra_gaps.get((frt_name, type_name), 0) + gap

for (frt_name, type_name), gap in extra_gaps.items():
    location = locations[frt_name]
    rtype = resource_types[type_name]
    env["gsus.requirement"].create({
        "location_id": location.id, "resource_type_id": rtype.id,
        "required_count": gap, "notes": "Additional gap scope from Sheet2",
    })
    env["gsus.frt.availability.line"].create({
        "availability_id": availability_by_frt[frt_name].id,
        "resource_type_id": rtype.id,
        "scope_count": gap, "available_count": 0,
        "remarks": "Resource not available - imported from Sheet2",
    })

# Preserve the specifically named unavailable Patrollers from Sheet2.
for row in range(12, sheet2.max_row + 1):
    description = str(sheet2.cell(row, 3).value or "").strip()
    frt_name = str(sheet2.cell(row, 4).value or "").strip()
    person_name = str(sheet2.cell(row, 5).value or "").strip()
    if description != "Patroller" or frt_name not in locations or person_name in ("", "-"):
        continue
    location = locations[frt_name]
    env["gsus.resource"].create({
        "name": person_name,
        "code": f"VI-GAP-{slug(frt_name)[:8]}-{slug(person_name)[:8]}",
        "resource_type_id": resource_types["Patroller"].id,
        "ownership": "gsus_own",
        "customer_id": customer.id,
        "project_id": project.id,
        "zone_id": location.zone_id.id,
        "location_id": location.id,
        "status": "unavailable",
        "shift_id": day_shift.id,
        "remarks": "Named unavailable resource from Sheet2",
    })

# One explicit leave/replacement example for client demonstration. The absent
# permanent helper is removed only from today's availability; master assignment
# remains unchanged. The temporary alternate is available for selection/testing.
demo_frt = locations.get("Purasaiwakkam")
if demo_frt:
    helper_line = availability_by_frt["Purasaiwakkam"].line_ids.filtered(
        lambda line: line.resource_type_id == resource_types["Helper"]
    )[:1]
    if helper_line and helper_line.available_resource_ids:
        absent = helper_line.available_resource_ids[-1]
        alternate = env["gsus.resource"].create({
            "name": "Purasaiwakkam Alternate Helper 01",
            "resource_type_id": resource_types["Helper"].id,
            "ownership": "vendor",
            "vendor_id": customer.id,
            "customer_id": customer.id,
            "project_id": project.id,
            "zone_id": demo_frt.zone_id.id,
            "location_id": demo_frt.id,
            "shift_id": day_shift.id,
            "employment_type": "temporary",
            "status": "available",
            "remarks": "Demo temporary replacement resource",
        })
        helper_line.write({
            "available_resource_ids": [(3, absent.id)],
            "alternate_resource_ids": [(6, 0, alternate.ids)],
            "available_count": len(helper_line.available_resource_ids) - 1,
            "alternate_count": 1,
            "remarks": "%s is Off; %s allocated as alternate" % (absent.name, alternate.name),
        })

env.cr.commit()
print("Vodafone demo created:", len(zones), "zones,", len(locations), "FRTs,", person_sequence - 1, "available persons,", vehicle_sequence - 1, "vehicles")
